import os
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
import pandas as pd

try:
  ROOT_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
except NameError:
  if os.path.basename(os.getcwd()) == "src":
    ROOT_DIR = os.path.dirname(os.getcwd())
  else:
    ROOT_DIR = os.getcwd()

os.chdir(ROOT_DIR)

# =====================================================================
# GLOBAL AYARLAR
# =====================================================================
EVALUATION = "outputs/evaluation.jsonl"
COMPANIES = "outputs/companies.json"
OUTPUT = "outputs/merged.xlsx"

# Excel'de görmek istediğin sütunlar ve sıralaması
SECILEN_SUTUNLAR = [
    "firma",
    "tanıtım",
    "link",
    "site_linki",
    "kategori",
    "ozet",
    "uygun_ilan_linkleri",
]

# Sütun adlarını Türkçeleştirmek veya değiştirmek istersen
SUTUN_YENIDEN_ADLANDIR = {
    "firma": "Firma Adı",
    "tanıtım": "Şirket Tanıtımı",
    "link": "Şirket Linki",
    "site_linki": "Web Sitesi",
    "kategori": "AI Kategorisi",
    "ozet": "AI Özeti",
    "uygun_ilan_linkleri": "Uygun İlan Linkleri",
}
# =====================================================================


def json_veya_jsonl_oku(dosya_yolu):
  if not os.path.exists(dosya_yolu):
    print(f"Uyarı: '{dosya_yolu}' dosyası bulunamadı, atlanıyor.")
    return None

  uzanti = os.path.splitext(dosya_yolu)[1].lower()

  if uzanti == ".jsonl":
    df = pd.read_json(dosya_yolu, lines=True)
  elif uzanti == ".json":
    try:
      df = pd.read_json(dosya_yolu)
    except ValueError:
      df = pd.read_json(dosya_yolu, lines=True)
  else:
    raise ValueError(f"Desteklenmeyen dosya formatı: {dosya_yolu}")

  df.columns = [str(col).strip().lower() for col in df.columns]
  return df


def main():
  print("Dosyalar okunuyor...")
  companies_df = json_veya_jsonl_oku(COMPANIES)
  evaluation_df = json_veya_jsonl_oku(EVALUATION)

  # 1. İç içe (nested) 'ai_analizi' sütununu düzleştirme
  if evaluation_df is not None and "ai_analizi" in evaluation_df.columns:
    ai_df = pd.json_normalize(evaluation_df["ai_analizi"])
    ai_df.columns = [str(col).strip().lower() for col in ai_df.columns]

    for col in ai_df.columns:
      ai_df[col] = ai_df[col].apply(
          lambda x: ", ".join(x) if isinstance(x, list) else x
      )

    evaluation_df = evaluation_df.drop(columns=["ai_analizi"])
    evaluation_df = pd.concat(
        [evaluation_df.reset_index(drop=True), ai_df.reset_index(drop=True)],
        axis=1,
    )

  # 2. İki veri setini 'firma' alanı üzerinden birleştirme
  if companies_df is not None and evaluation_df is not None:
    companies_df["firma"] = companies_df["firma"].astype(str).str.strip()
    evaluation_df["firma"] = evaluation_df["firma"].astype(str).str.strip()
    df = pd.merge(companies_df, evaluation_df, on="firma", how="outer")
  elif companies_df is not None:
    df = companies_df
  elif evaluation_df is not None:
    df = evaluation_df
  else:
    print("Hata: Hiçbir veri dosyası okunamadı!")
    return

  # Liste sütunlarını metne çevir
  for col in df.columns:
    df[col] = df[col].apply(
        lambda x: ", ".join(x) if isinstance(x, list) else x
    )

  # 3. KESİN ALFABETİK SIRALAMA (Büyük/küçük harf duyarsız A-Z)
  if "firma" in df.columns:
    df = df.sort_values(
        by="firma", key=lambda col: col.str.lower(), ascending=True
    ).reset_index(drop=True)

  # 4. Sütun filtreleme
  if SECILEN_SUTUNLAR:
    mevcut_sutunlar = [s for s in SECILEN_SUTUNLAR if s in df.columns]
    df = df[mevcut_sutunlar]

  # 5. Sütun isimlerini değiştirme (Türkçeleştirme)
  if SUTUN_YENIDEN_ADLANDIR:
    df = df.rename(columns=SUTUN_YENIDEN_ADLANDIR)

  # 6. Klasör kontrolü ve Excel'e kaydetme
  cikis_klasoru = os.path.dirname(OUTPUT)
  if cikis_klasoru and not os.path.exists(cikis_klasoru):
    os.makedirs(cikis_klasoru)

  df.to_excel(OUTPUT, index=False, engine="openpyxl")

  # =====================================================================
  # 7. OKUNABİLİRLİK VE PROFESYONEL EXCEL BİÇİMLENDİRME (OPENPYXL)
  # =====================================================================
  wb = openpyxl.load_workbook(OUTPUT)
  ws = wb.active
  ws.title = "Firma Analizleri"

  # Başlık satırını sabitle ve filtre ekle
  ws.freeze_panes = "A2"
  if ws.dimensions:
    ws.auto_filter.ref = ws.dimensions

  # Renk ve Stil Tanımları (Times New Roman olarak güncellendi)
  header_font = Font(
      name="Times New Roman", size=11, bold=True, color="FFFFFF"
  )
  header_fill = PatternFill(
      start_color="2C3E50", end_color="2C3E50", fill_type="solid"
  )
  header_alignment = Alignment(
      horizontal="center", vertical="center", wrap_text=True
  )

  cell_font = Font(name="Times New Roman", size=11)
  cell_alignment = Alignment(vertical="top", horizontal="left", wrap_text=True)

  thin_border = Border(
      left=Side(style="thin", color="E0E0E0"),
      right=Side(style="thin", color="E0E0E0"),
      top=Side(style="thin", color="E0E0E0"),
      bottom=Side(style="thin", color="E0E0E0"),
  )

  zebra_fill = PatternFill(
      start_color="F8F9F9", end_color="F8F9F9", fill_type="solid"
  )

  # Başlıkları Biçimlendir
  ws.row_dimensions[1].height = 28
  for col_num in range(1, ws.max_column + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

  # Veri Satırlarını Biçimlendir
  for row_num in range(2, ws.max_row + 1):
    ws.row_dimensions[row_num].height = 45
    is_even = row_num % 2 == 0
    for col_num in range(1, ws.max_column + 1):
      cell = ws.cell(row=row_num, column=col_num)
      cell.font = cell_font
      cell.alignment = cell_alignment
      cell.border = thin_border
      if is_even:
        cell.fill = zebra_fill

  # Sütun Genişliklerini İçeriğe Göre Otomatik Ayarla
  for col in ws.columns:
    max_len = 0
    col_letter = get_column_letter(col[0].column)
    for cell in col:
      val = str(cell.value or "")
      line_len = min(len(val), 45)
      if line_len > max_len:
        max_len = line_len
    ws.column_dimensions[col_letter].width = max(max_len + 6, 18)

  wb.save(OUTPUT)
  print(
      f"Başarılı! Times New Roman formatlı ve alfabetik sıralı veriler"
      f" '{OUTPUT}' dosyasına kaydedildi."
  )


if __name__ == "__main__":
  main()